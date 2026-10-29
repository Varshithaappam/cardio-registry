import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOWNLOADS_DIR = r"C:\Users\ADMIN\Downloads"
EXCEL_PATH = os.path.join(DOWNLOADS_DIR, "DDH-DR-CR3 - UAT_Release_Notes_Action_List (20261029).xlsx")
DOCX_PATH = os.path.join(DOWNLOADS_DIR, "DDH-DR-CR3 - UAT_Release_Notes (20261029).docx")

# ==============================================================================
# 1. ACTION LIST DATA (Detailed changes after CR2)
# ==============================================================================
ACTION_ITEMS = [
    {
        "id": "CR3-001",
        "module": "Clinical Intelligence Dashboard",
        "feature": "Heart Failure GDMT Clinical Intelligence Engine",
        "description": "Architected and deployed full-stack GDMT adherence calculation engine evaluating the 4 core guideline pillars (Beta-Blocker, ACEi/ARB/ARNI, MRA, SGLT2i) in real-time across patient cohorts.",
        "type": "New Feature",
        "files": "frontend/src/components/Dashboard.jsx, backend/services/gdmtService.js, frontend/src/utils/gdmtEngine.js",
        "impact": "High - Provides real-time clinical visibility into guideline adherence and therapy gaps.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-002",
        "module": "Clinical Intelligence Dashboard",
        "feature": "Dynamic Contraindication Detection System",
        "description": "Built automated detection rules for physiological and clinical contraindications (hypotension SBP < 85 mmHg, hyperkalemia K+ > 5.5 mEq/L, severe renal impairment Creatinine > 2.5 mg/dL, cardiogenic shock, bradycardia, angioedema) to distinguish between justified omissions and true therapy deviations.",
        "type": "New Feature",
        "files": "backend/services/gdmtEvaluationService.js, frontend/src/components/dashboard/ContraindicationDrawer.jsx",
        "impact": "High - Prevents incorrect deviation flags when medication is clinically contraindicated.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-003",
        "module": "Clinical Intelligence Dashboard",
        "feature": "4-Panel Clinical Epidemiology Grid",
        "description": "Implemented high-impact visual epidemiology panels: (1) Gender Distribution donut chart, (2) HF Phenotype Classification (HFrEF vs. HFpEF donut chart), (3) NYHA Functional Class semi-circle radial gauge, and (4) Age Distribution Spectrum with peak density curve.",
        "type": "New Feature",
        "files": "frontend/src/components/dashboard/KeyPointsTab.jsx",
        "impact": "High - Delivers immediate demographic and diagnostic profiling of the registered HF cohort.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-004",
        "module": "Clinical Intelligence Dashboard",
        "feature": "Universal Cohort Filter Bar",
        "description": "Created interactive global filter bar enabling real-time cohort slicing across Registry type, HF Phenotype (HFrEF/HFpEF), Gender, Age range, Admission date range, and free-text search (UHID/Name/Drugs).",
        "type": "New Feature",
        "files": "frontend/src/components/dashboard/UniversalFilterBar.jsx",
        "impact": "Medium - Allows clinical research teams to drill down into specific sub-cohorts instantly.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-005",
        "module": "Clinical Intelligence Dashboard",
        "feature": "Patient Dossier & Interactive Contraindication Drawer",
        "description": "Integrated interactive slide-out drawers and dossier modals allowing clinicians to click any patient or metric to inspect complete drug profiles, contraindication rationales, and visit history.",
        "type": "New Feature",
        "files": "frontend/src/components/dashboard/PatientDossierModal.jsx, ContraindicationDrawer.jsx",
        "impact": "Medium - Provides transparency behind GDMT adherence scores and deviation calculations.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-006",
        "module": "Clinical Intelligence Dashboard",
        "feature": "Dashboard Label Auto-Sizing & Truncation Elimination",
        "description": "Eliminated text truncation (ellipsis clipping) across all dashboard cards. Restructured badges into responsive vertical layouts with auto-height (min-h-[58px]) so NYHA Class 3 (Class III), Class 4 (Class IV), HFrEF, HFpEF, Female, and multi-line titles are 100% visible.",
        "type": "Bug Fix / UX",
        "files": "frontend/src/components/dashboard/KeyPointsTab.jsx",
        "impact": "Medium - Completely resolves UI truncation where Class 3, Class 4, and HFrEF were hidden as 'Class...' and 'HF...'.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-007",
        "module": "Clinical Intelligence Dashboard",
        "feature": "Comorbidity Burden & Mortality Endpoints Analytics",
        "description": "Built dedicated analytics for Diabetes Mellitus, Hypertension, and Renal Failure/CKD comorbidity rates, alongside in-hospital, post-discharge, and 30-day timeline mortality tracking.",
        "type": "Enhancement",
        "files": "frontend/src/components/dashboard/KeyPointsTab.jsx, frontend/src/utils/gdmtEngine.js",
        "impact": "Medium - Gives cardiology teams actionable insights into high-risk clinical comorbidities.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-008",
        "module": "Excel Export Engine",
        "feature": "Institutional Multi-Sheet Excel Workbook Export",
        "description": "Developed professional ExcelJS export engine generating a 5-sheet audit workbook: (1) Key Points Executive Summary, (2) GDMT Summary, (3) Sheet 1 Master Clinical Cohort Dataset, (4) Notes & Clinical Definitions, (5) Contraindications Matrix.",
        "type": "New Feature",
        "files": "frontend/src/utils/excelReportGenerator.js, backend/services/hfExportService.js",
        "impact": "High - Allows automated generation of hospital registry submission files and clinical audit sheets.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-009",
        "module": "Excel Export Engine",
        "feature": "Generic Drug Names Aggregation (Column BU)",
        "description": "Implemented automated extraction and compilation of all generic drug prescriptions across medical therapy parts 1, 2, and 3 into Column BU where a selection, tick, or dosage is recorded.",
        "type": "Enhancement",
        "files": "backend/services/hfExportService.js, frontend/src/utils/excelReportGenerator.js",
        "impact": "High - Ensures comprehensive pharmacological record without manual data entry.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-010",
        "module": "Excel Export Engine",
        "feature": "Prescription-Driven GDMT Pillar Classification",
        "description": "Refactored GDMT pillar categorization in the export file to dynamically specify the exact drug molecule prescribed (e.g., 'Beta Blocker (Metoprolol Succinate)') rather than static hardcoded category strings.",
        "type": "Enhancement",
        "files": "backend/services/hfExportService.js",
        "impact": "Medium - Provides clinical clarity on which specific drug satisfies each GDMT guideline pillar.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-011",
        "module": "Excel Export Engine",
        "feature": "ECG Conduction Blockages Dynamic Resolution",
        "description": "Fixed export mapping for heart blockages so that any positive conduction abnormality across ecg_lbbb, ecg_rbbb, or ecg_block_other dynamically displays its exact diagnostic name (LBBB, RBBB, or specified block).",
        "type": "Bug Fix / Data Quality",
        "files": "backend/services/hfExportService.js, frontend/src/utils/excelReportGenerator.js",
        "impact": "Medium - Corrects diagnostic blockage reporting in the institutional export file.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-012",
        "module": "Excel Export Engine",
        "feature": "Export Column Formatting & Header Standardization",
        "description": "Removed redundant 'Reason for Admission 2' column, renamed 'Reason for Admission 3' to 'Admission Type', and standardized Column 2 (AF) header to 'O₂ Saturation (%)' with proper numeric precision.",
        "type": "Enhancement",
        "files": "frontend/src/utils/excelReportGenerator.js, backend/services/hfExportService.js",
        "impact": "Medium - Aligns Excel export columns strictly with hospital registry guidelines.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-013",
        "module": "Excel Export Engine",
        "feature": "Export UHID Filtering & Data Inclusion Fix",
        "description": "Identified and resolved export exclusion where patients lacking explicit UHID values in patient_demographics were being filtered out, ensuring all registered HF records can be successfully audited and exported.",
        "type": "Bug Fix / Data Quality",
        "files": "backend/services/hfExportService.js",
        "impact": "High - Prevents registered patients from being omitted from administrative export files.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-014",
        "module": "Data Quality & Integrity",
        "feature": "Centralized Patient Demographics & UHID Alignment",
        "description": "Standardized patient linking between dbo.hf_registry and dbo.patient_demographics by reg_patient_id, preventing orphan records and ensuring verified UHIDs propagate accurately to all clinical forms.",
        "type": "Enhancement",
        "files": "backend/models/patientModel.js, backend/services/hfService.js",
        "impact": "High - Preserves relational integrity across patient registration and disease registry forms.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-015",
        "module": "Data Quality & Integrity",
        "feature": "Child Table Redundancy Elimination",
        "description": "Audited and resolved duplicate row creation across HF child tables (hf_administrative, hf_initial_assessment, hf_final_clinical_assessment, etc.) caused by multiple form updates.",
        "type": "Data Quality",
        "files": "backend/services/hfService.js, backend/migrations/ddl_changes.sql",
        "impact": "High - Guarantees 1-to-1 integrity between primary registry entries and child clinical sections.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-016",
        "module": "Registry Clinical Forms",
        "feature": "Removal of Artificial Default Form Values",
        "description": "Purged pre-selected default / placeholder values across HF, STEMI, and NSTEMI clinical forms, ensuring that unassessed clinical parameters remain strictly null/blank rather than skewing data.",
        "type": "Bug Fix / Data Quality",
        "files": "frontend/src/components/forms/HFForm.jsx, frontend/src/components/forms/STEMIForm.jsx, NSTEMIForm.jsx",
        "impact": "High - Ensures high-fidelity clinical data entry without biased default assumptions.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-017",
        "module": "Registry Clinical Forms",
        "feature": "Automated Form Validation Error Navigation",
        "description": "Implemented fieldScrollHelper and ScrollToTop utilities to automatically scroll the screen and highlight invalid or missing required fields upon form submission, preventing user frustration on long forms.",
        "type": "UX Enhancement",
        "files": "frontend/src/utils/fieldScrollHelper.js, frontend/src/components/ScrollToTop.jsx",
        "impact": "Medium - Drastically improves nurse and clinician workflow efficiency during data submission.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-018",
        "module": "Registry Clinical Forms",
        "feature": "Typography, Spacing & Layout Stabilization",
        "description": "Standardized font sizes, section card headers, medication dosage rows, table alignments, and responsive input styles across clinical assessment views.",
        "type": "UX Enhancement",
        "files": "frontend/src/components/forms/common/formStyles.js, frontend/src/components/ClinicalForm.jsx",
        "impact": "Low - Polishes user interface aesthetics and readability on medical workstations.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-019",
        "module": "Nurse Follow-Up Workflow",
        "feature": "Nurse Queue Workflow & ACS Follow-Up PDF Upgrades",
        "description": "Enhanced Nurse Follow-up queue with multi-registry filtering, auto-transitioning completed outreach to 'Detailed Form Logged', and upgraded ACS Follow-up PDF report modal with visit mode and clinical notes.",
        "type": "Enhancement",
        "files": "frontend/src/components/NurseFollowUpReport.jsx, frontend/src/components/modals/AcsFollowupPdfModal.jsx",
        "impact": "Medium - Streamlines post-discharge patient outreach tracking for nurse coordinators.",
        "status": "Completed & Verified"
    },
    {
        "id": "CR3-020",
        "module": "Backend & Database",
        "feature": "Database Connection Pooling & Query Performance",
        "description": "Optimized MSSQL database queries for cohort aggregation and export retrieval, implemented resilient query timeout configurations, and added server startup health check verification.",
        "type": "Performance",
        "files": "backend/config/db.js, backend/services/gdmtService.js",
        "impact": "High - Prevents connection pool exhaustion and database lockups during bulk cohort processing.",
        "status": "Completed & Verified"
    }
]

# ==============================================================================
# 2. GENERATE EXCEL ACTION LIST WORKBOOK
# ==============================================================================
def generate_excel():
    wb = openpyxl.Workbook()

    # Sheet 1: Action List Table
    ws1 = wb.active
    ws1.title = "CR3 Action List"
    ws1.views.sheetView[0].showGridLines = True

    # Palette
    NAVY = "1B365D"
    TEAL = "008080"
    LIGHT_TEAL = "E6F4F1"
    ZEBRA_FILL = "F8FAFC"
    WHITE = "FFFFFF"
    BORDER_COLOR = "CBD5E1"
    GRAY_TEXT = "475569"

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    # Title block
    ws1.merge_cells("A1:H1")
    title_cell = ws1["A1"]
    title_cell.value = "CARE HEALTH SYSTEM - CARDIOLOGY DISEASE REGISTRY"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color=WHITE)
    title_cell.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:H2")
    sub_cell = ws1["A2"]
    sub_cell.value = "Change Request 3 (CR3) Action List & UAT Release Modifications (Post-CR2 to 29th October 2026)"
    sub_cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    sub_cell.fill = PatternFill(start_color=TEAL, end_color=TEAL, fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 24

    # Metadata row
    ws1.merge_cells("A3:H3")
    meta_cell = ws1["A3"]
    meta_cell.value = "CR Number: DDH-DR-CR3   |   Target Environment: UAT   |   Planned Deployment Date: 29th October 2026   |   Baseline Release: CR2 (20261006)"
    meta_cell.font = Font(name="Calibri", size=10, italic=True, color=GRAY_TEXT)
    meta_cell.fill = PatternFill(start_color=LIGHT_TEAL, end_color=LIGHT_TEAL, fill_type="solid")
    meta_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[3].height = 22

    # Empty spacer
    ws1.row_dimensions[4].height = 10

    # Headers
    headers = [
        "Item ID",
        "Module / Subsystem",
        "Feature / Action Item",
        "Detailed Change Description & Technical Scope",
        "Change Type",
        "Primary Files Modified",
        "Clinical & Operational Impact",
        "UAT Status"
    ]

    header_row = 5
    ws1.row_dimensions[header_row].height = 28
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws1.cell(row=header_row, column=col_idx)
        cell.value = header_text
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Populate Action Items
    current_row = 6
    for item in ACTION_ITEMS:
        ws1.row_dimensions[current_row].height = 48
        is_even = (current_row % 2 == 0)
        row_fill = PatternFill(start_color=ZEBRA_FILL if is_even else WHITE,
                               end_color=ZEBRA_FILL if is_even else WHITE, fill_type="solid")

        cells = [
            (1, item["id"], "center", Font(name="Calibri", size=10, bold=True, color=NAVY)),
            (2, item["module"], "center", Font(name="Calibri", size=10, bold=True)),
            (3, item["feature"], "left", Font(name="Calibri", size=10, bold=True)),
            (4, item["description"], "left", Font(name="Calibri", size=9.5)),
            (5, item["type"], "center", Font(name="Calibri", size=10, bold=True)),
            (6, item["files"], "left", Font(name="Consolas", size=8.5, color=GRAY_TEXT)),
            (7, item["impact"], "left", Font(name="Calibri", size=9.5)),
            (8, item["status"], "center", Font(name="Calibri", size=10, bold=True, color="00703C"))
        ]

        for col_idx, val, align_h, font_style in cells:
            cell = ws1.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_style
            cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=True)
            cell.border = thin_border
            cell.fill = row_fill

        current_row += 1

    # Column widths
    col_widths = {
        1: 12,  # Item ID
        2: 24,  # Module
        3: 28,  # Feature
        4: 48,  # Description
        5: 18,  # Type
        6: 32,  # Files
        7: 34,  # Impact
        8: 18   # Status
    }
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = width

    # Sheet 2: Executive Release Summary
    ws2 = wb.create_sheet(title="Executive Summary")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:F1")
    t2 = ws2["A1"]
    t2.value = "RELEASE NOTES: DISEASE REGISTRY IN UAT ENVIRONMENT (CR3)"
    t2.font = Font(name="Calibri", size=14, bold=True, color=WHITE)
    t2.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    t2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 32

    summary_rows = [
        ("Change Request No.", "DDH-DR-CR3"),
        ("Schedule Planned Date", "29th October 2026"),
        ("Schedule Time Window", "2:00 PM - 5:00 PM IST"),
        ("Baseline Prior Release", "DDH-DR-CR2 (Deployed 6th October 2026)"),
        ("Environment", "User Acceptance Testing (UAT)"),
        ("Primary Business Owner", "Cardiology Disease Registry & Clinical Research Department")
    ]

    for idx, (label, val) in enumerate(summary_rows, start=3):
        ws2.row_dimensions[idx].height = 22
        c_lbl = ws2.cell(row=idx, column=1, value=label)
        c_lbl.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        c_lbl.fill = PatternFill(start_color=LIGHT_TEAL, end_color=LIGHT_TEAL, fill_type="solid")
        c_lbl.border = thin_border

        ws2.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=6)
        c_val = ws2.cell(row=idx, column=2, value=val)
        c_val.font = Font(name="Calibri", size=10)
        c_val.border = thin_border

    # Scope of Change Block
    r = 10
    ws2.merge_cells(f"A{r}:F{r}")
    sc_hdr = ws2.cell(row=r, column=1, value="Core Pillars & Scope of Change for CR3")
    sc_hdr.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    sc_hdr.fill = PatternFill(start_color=TEAL, end_color=TEAL, fill_type="solid")
    ws2.row_dimensions[r].height = 24
    r += 1

    scope_points = [
        "1. Heart Failure Clinical Intelligence & GDMT Adherence Dashboard: Real-time adherence evaluation across 4 pillars, automated contraindication detection, 4-panel epidemiology visual grid, and comorbidity burden analysis.",
        "2. Institutional Multi-Sheet Excel Export Engine: Automated generation of 5-sheet audit workbooks with dynamic generic drug extraction (Column BU), prescription-driven GDMT pillar labeling, and blockage resolution.",
        "3. Dashboard UI Auto-Sizing & Truncation Elimination: Restructured all metric cards with auto-adjusting heights (min-h-[58px]) so NYHA Class 3, Class 4, HFrEF, and multi-line titles are fully readable without ellipsis.",
        "4. Data Quality & Relational Integrity: Prevention of duplicate child table rows, elimination of artificial default input values, and strict synchronization between master patient index UHIDs and disease registry records.",
        "5. Clinical Form Usability & Error Navigation: Field scroll helper and scroll-to-top integrations to automatically focus and highlight missing required inputs on lengthy forms.",
        "6. Nurse Follow-up Workflow & ACS PDF Enhancements: Enhanced task queue filters, outreach tracking, and verified ACS follow-up PDF reporting.",
        "7. Database Performance & Query Optimization: Resilient connection pooling, optimized aggregation queries, and startup schema integrity checks."
    ]

    for pt in scope_points:
        ws2.merge_cells(f"A{r}:F{r}")
        cell = ws2.cell(row=r, column=1, value=pt)
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = thin_border
        ws2.row_dimensions[r].height = 36
        r += 1

    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 20
    ws2.column_dimensions["D"].width = 20
    ws2.column_dimensions["E"].width = 20
    ws2.column_dimensions["F"].width = 20

    wb.save(EXCEL_PATH)
    print("Successfully generated Excel:", EXCEL_PATH)


# ==============================================================================
# 3. GENERATE WORD RELEASE NOTES DOCUMENT (Identical format to CR2)
# ==============================================================================
def generate_word():
    doc = docx.Document()

    # Set margins to 1 inch
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Helper for adding paragraph
    def add_p(text, bold=False, italic=False, size_pt=11, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Calibri"
        run.font.size = Pt(size_pt)
        if color:
            run.font.color.rgb = color
        return p

    def add_bullet(text, bold_prefix=None, size_pt=11, space_after=4):
        p = doc.add_paragraph(style='List Paragraph')
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r1 = p.add_run(bold_prefix)
            r1.bold = True
            r1.font.name = "Calibri"
            r1.font.size = Pt(size_pt)
        r2 = p.add_run(text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(size_pt)
        return p

    def add_section_header(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return p

    # 1. Title Block
    add_p("Release Notes: Disease Registry in UAT environment", bold=True, size_pt=16, color=RGBColor(0x1B, 0x36, 0x5D), space_after=8)
    
    # Metadata Block
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    p_meta.paragraph_format.line_spacing = 1.15
    r = p_meta.add_run("Change Request No.: ")
    r.bold = True
    r.font.name = "Calibri"
    r = p_meta.add_run("DDH-DR-CR3\n")
    r.font.name = "Calibri"

    r = p_meta.add_run("Schedule Change Planned Date: ")
    r.bold = True
    r.font.name = "Calibri"
    r = p_meta.add_run("29th October 2026\n")
    r.font.name = "Calibri"

    r = p_meta.add_run("Schedule Time: ")
    r.bold = True
    r.font.name = "Calibri"
    r = p_meta.add_run("2pm-5pm")
    r.font.name = "Calibri"

    # 2. Scope of the Change
    add_section_header("Scope of the Change")
    add_p("This release brings major functional additions, clinical analytics, institutional export capabilities, and extensive data quality stabilization across the Disease Registry application following CR2. The key updates include:", bold=True, space_after=6)
    add_bullet("Deployment of the Heart Failure Clinical Intelligence & GDMT Adherence Dashboard.")
    add_bullet("Dynamic Contraindication Detection Engine distinguishing justified clinical omissions from true guideline deviations.")
    add_bullet("Institutional Master Excel Export Engine with strict mapping, dynamic drug extraction, and conduction blockage resolution.")
    add_bullet("Dashboard UI label auto-sizing and removal of text truncation across NYHA classes, HF phenotypes, and card headers.")
    add_bullet("Elimination of duplicate child table rows and relational alignment between master patient index UHIDs and disease registry records.")
    add_bullet("Purging of artificial pre-selected default form values across HF, STEMI, and NSTEMI clinical assessment forms.")
    add_bullet("Form submission error scrolling helper with automatic focus and visual highlighting.")
    add_bullet("Nurse Follow-Up task queue upgrades and enhanced ACS Follow-up PDF report generation.")
    add_bullet("Database query optimizations, connection pooling resilience, and automated schema synchronization.")

    # 3. Section 1: HF Clinical Intelligence & GDMT Dashboard
    add_section_header("Heart Failure Clinical Intelligence & GDMT Adherence Dashboard")
    add_p("A comprehensive clinical intelligence dashboard has been implemented for the Heart Failure registry (/dashboard) providing deep analytics into guideline-directed medical therapy. This includes:")
    add_bullet("4 GDMT Therapy Pillars: Evaluates patient adherence across Beta-Blockers, ACEi/ARB/ARNI, MRA, and SGLT2 Inhibitors.")
    add_bullet("Dynamic Contraindication Detection: Real-time clinical evaluation of contraindications including hypotension (SBP < 85 mmHg), hyperkalemia (K+ > 5.5 mEq/L), severe renal impairment (Creatinine > 2.5 mg/dL), cardiogenic shock, bradycardia, and angioedema.")
    add_bullet("4-Panel Clinical Epidemiology Grid: High-level visual breakdown of Gender Distribution, HF Phenotype Classification (HFrEF vs. HFpEF), NYHA Functional Class (Radial gauge), and Age Distribution Spectrum.")
    add_bullet("Comorbidity Burden Analysis: Clinical tracking and percentage prevalence of Diabetes Mellitus, Hypertension, and Renal Failure / CKD.")
    add_bullet("Mortality Endpoints & Timeline: Detailed quantification of in-hospital deaths, post-discharge mortality, and 30-day post-discharge survival timeline.")
    add_bullet("Universal Cohort Filter Bar: Dynamic cohort slicing by Registry, HF Phenotype, Gender, Age bracket, Admission date range, and free-text patient search.")
    add_bullet("Patient Dossier & Interactive Contraindication Drawer: Slide-out clinical audit drawers displaying individual drug dosages, adherence breakdowns, and clinical rationale.")

    # 4. Section 2: Dashboard UI Auto-Sizing & Truncation Elimination
    add_section_header("Dashboard UI Auto-Sizing & Label Truncation Elimination")
    add_p("Dashboard metric cards and legend badges have been overhauled to dynamically size their dimensions based on inner content, completely removing text clipping:")
    add_bullet("NYHA Functional Class (Class 3 & Class 4): Removed text truncation ('Class...') and added dual labeling (Class III / Class 3, Class IV / Class 4) so clinical classes are 100% visible.")
    add_bullet("HF Phenotype Classification: Replaced clipped 'HF...' with full 'HFrEF' and 'HFpEF' badges.")
    add_bullet("Auto-Height Flex Layout: Replaced rigid horizontal constraints with vertical flex containers (min-h-[58px]) that naturally grow when text wraps.")
    add_bullet("Responsive Panel Headers: Standardized panel titles with multi-line word wrapping and fixed-badge alignment to prevent vertical compression.")

    # 5. Section 3: Institutional Master Excel Export Engine
    add_section_header("Institutional Multi-Sheet Excel Export Engine")
    add_p("A hospital-grade multi-sheet Excel export engine has been deployed (excelReportGenerator.js & hfExportService.js) strictly matching institutional registry reporting guidelines:")
    add_bullet("Multi-Sheet Audit Workbook: Automatically exports 5 coordinated sheets: Key Points Executive Summary, GDMT Summary, Sheet 1 (Master HF Cohort Dataset), Notes, and Contraindications.")
    add_bullet("Generic Drug Names Aggregation (Column BU): Dynamically aggregates all prescribed generic drug names across medical therapy parts 1, 2, and 3 where doses or selections are recorded.")
    add_bullet("Prescription-Driven GDMT Pillar Classification: Dynamically maps the exact drug molecule taken (e.g. 'Beta Blocker (Metoprolol Succinate)') rather than generic static pillar labels.")
    add_bullet("Conduction Blockage Resolution: Automatically displays the specific blockage name ('LBBB', 'RBBB', or specific custom block) based on ECG diagnostic flags.")
    add_bullet("Column Header Standardization: Removed redundant 'Reason for Admission 2' column, renamed 'Reason for Admission 3' to 'Admission Type', and standardized Column 2 to 'O₂ Saturation (%)'.")
    add_bullet("UHID Alignment & Export Query Fix: Resolved query filtering so all active registered HF patients export reliably regardless of provisional UHID status.")

    # 6. Section 4: Data Quality, Registry Integrity & Form UX
    add_section_header("Data Quality, Form Stabilization & Relational Integrity")
    add_p("Clinical data collection forms and database relational integrity have been standardized:")
    add_bullet("Centralized Demographics & UHID Integration: Synchronized HF registry records with master patient index UHIDs, preventing orphan records and duplicate registrations.")
    add_bullet("Child Table Redundancy Elimination: Audited and prevented multi-row insertions in HF clinical child tables (hf_administrative, hf_initial_assessment, hf_final_clinical_assessment, etc.).")
    add_bullet("Removal of Artificial Default Form Values: Purged pre-selected default values across HF, STEMI, and NSTEMI forms so unassessed fields remain strictly null rather than skewing clinical statistics.")
    add_bullet("Automated Validation Error Navigation: Integrated fieldScrollHelper and ScrollToTop to smoothly scroll and highlight missing or invalid inputs during submission.")
    add_bullet("Typography & Spacing Harmonization: Unified clinical font sizes, medication row padding, and input styling across all registry sections.")

    # 7. Section 5: Nurse Follow-Up & ACS PDF Enhancements
    add_section_header("Nurse Follow-Up Workflow & ACS PDF Enhancements")
    add_p("The Nurse Follow-up module has received functional upgrades to streamline post-discharge tracking:")
    add_bullet("Task Queue Workflow: Multi-registry filtering across HF, STEMI, and NSTEMI queues with automatic status transitions to 'Detailed Form Logged'.")
    add_bullet("ACS Follow-Up PDF Report: Upgraded PDF modal for STEMI and NSTEMI follow-ups capturing outreach notes, visit mode (In-person vs. Telephonic), and clinical recommendations.")

    # 8. Section 6: Database & Backend Architecture
    add_section_header("Database Architecture & Performance Optimization")
    add_p("Backend database handling has been fortified for heavy clinical queries:")
    add_bullet("High-Performance Cohort Evaluation: Deployed backend evaluation services (gdmtService.js, gdmtEvaluationService.js) for high-speed cohort analytics.")
    add_bullet("MSSQL Connection Resilience: Reconfigured connection pooling and query timeouts to prevent connection drops during bulk exports.")
    add_bullet("Automated Schema Synchronization: Verified table constraints, audit log triggers, and foreign keys across all registry tables.")

    # 9. Testing & Deployment Validation
    add_section_header("Testing & Deployment Validation")
    add_p("Immediately after UAT deployment, a smoke test will be performed to confirm that the environment is functioning correctly. The following will be validated:")
    add_bullet("User login and authentication.")
    add_bullet("Clinical Intelligence Dashboard loading, filter responsiveness, and KPI calculations.")
    add_bullet("GDMT adherence evaluation and contraindication drawer display.")
    add_bullet("Multi-sheet Excel export generation, generic drugs extraction (Column BU), and blockage mapping.")
    add_bullet("Clinical form validation error scrolling and submission without default value pollution.")
    add_bullet("Nurse Follow-up task queue transitions and ACS PDF report generation.")
    add_bullet("Database connection pool stability under concurrent queries.")
    add_bullet("Application logs and error monitoring.")
    add_p("A 5-day hypercare period will be provided following deployment to closely monitor system performance, address any critical issues, support users, and ensure the stability of the new functionality before transitioning to regular support.")

    # 10. Rollback Plan
    add_section_header("Rollback Plan")
    add_p("In case of any critical issue following deployment, the release can be rolled back to the previous stable application version (CR2):")
    add_bullet("Application rollback: The previous stable application version (CR2 commit e182f1d) will be restored if critical issues arise.")
    add_bullet("Data protection: A complete database backup will be taken immediately prior to deployment to ensure all patient and clinical data can be recovered.")
    add_bullet("Audit records: System audit logs, clinical change logs, and patient records will be fully preserved.")
    add_bullet("Database changes: Database rollback will only be performed if strictly necessary after verifying data integrity.")
    add_bullet("Post-rollback validation: Core patient registration, HF/STEMI/NSTEMI workflows, and follow-up activities will be validated prior to reopening system access.")
    add_p("Rollback will prioritize system stability while ensuring that no patient or clinical registry data is lost.")

    # 11. Release Sign-Off
    add_section_header("Release Sign-Off")
    add_p("The release will be considered successfully deployed once:")
    add_bullet("All critical clinical workflows and dashboard analytics function as expected.")
    add_bullet("No critical or high-severity defects are identified during smoke testing.")
    add_bullet("Patient and registry data integrity is fully confirmed.")
    add_bullet("Key user journeys (Dashboard, Forms, Nurse Follow-Up, Excel Export) pass validation.")
    add_bullet("The Product and Technical teams provide final release sign-off.")

    doc.save(DOCX_PATH)
    print("Successfully generated Word Doc:", DOCX_PATH)

if __name__ == "__main__":
    generate_excel()
    generate_word()
